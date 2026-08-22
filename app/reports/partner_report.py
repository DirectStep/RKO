import io
from datetime import date
from uuid import UUID

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from app.database import Database
from app.services.partner_cabinet import PartnerBankData, partner_cabinet_data

LEAD_STATUS_LABELS = {
    "new": "Новая",
    "in_progress": "В работе",
    "opening_accounts": "Открытие счетов",
    "partially_completed": "Частично завершена",
    "completed": "Завершена",
    "paused": "На паузе",
    "closed_without_result": "Закрыта без результата",
}
BANK_STATUS_LABELS = {
    "planned": "Запланирован",
    "in_progress": "В работе",
    "opened": "Открыт",
    "not_opened": "Не открыт",
    "will_not_open": "Не будет открыт",
}
PAYMENT_STATUS_LABELS = {
    "not_calculated": "Не рассчитана",
    "calculated": "Рассчитана",
    "awaiting_confirmation": "Ждёт подтверждения",
    "confirmed": "Подтверждена",
    "in_registry": "В реестре",
    "paid": "Выплачена",
    "cancelled": "Отменена",
}


async def build_partner_report(
    database: Database,
    partner_id: UUID,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    data = await partner_cabinet_data(
        database,
        partner_id,
        date_from=date_from,
        date_to=date_to,
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заявки"
    headers = [
        "Заявка",
        "Клиент",
        "Telegram",
        "Дата",
        "Канал",
        "Статус заявки",
        "Банк",
        "Статус банка",
        "Расчётная выплата",
        "Подтверждённая выплата",
        "Статус выплаты",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="229ED9")
    for lead in data["leads"]:
        banks: list[PartnerBankData | None] = [*lead["banks"]]
        if not banks:
            banks.append(None)
        lead_status = LEAD_STATUS_LABELS[lead["status"]]
        if lead["is_repeat"]:
            lead_status = f"Повторная · {lead_status}"
        for bank in banks:
            sheet.append(
                [
                    lead["short_id"],
                    lead["name"],
                    lead["username"],
                    str(lead["date"])[:10],
                    lead["channel"],
                    lead_status,
                    bank["bank"] if bank else "",
                    BANK_STATUS_LABELS[bank["status"]] if bank else "",
                    float(bank["reward_estimate"]) if bank else 0,
                    float(bank["reward_fact"]) if bank else 0,
                    PAYMENT_STATUS_LABELS[bank["payment_status"]] if bank else "",
                ]
            )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in enumerate((15, 26, 20, 13, 22, 24, 20, 20, 20, 24, 22), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2, min_col=9, max_col=10):
        for cell in row:
            cell.number_format = '#,##0.00 "₽"'
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
